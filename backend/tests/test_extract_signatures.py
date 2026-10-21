"""The attack signatures of ZIP files and Office packages (D-36-17, 36-WORTLAUT.md A3, plan 36-08).

Pure tests in this process, without the sandbox: the corpus of
test_phase36_hostile_corpus.py runs the same families through the real child.
Here every edge of every rule stands on its own:

* **zip_ratio, two steps (answer O12).** A member from ARCHIVE_RATIO_MIN_BYTES
  on with a ratio over ARCHIVE_MAX_RATIO (100) up to and including
  ARCHIVE_MALICIOUS_RATIO (1000) stays archive_limit; only above that number
  is the file blocked as malicious. The ratios are built through lying
  declarations of a stored member where an exact number is needed, and every
  test reads the ratio back from the bytes (file_size over compress_size)
  instead of believing the generator.
* **zip_overlap on every level (answer O2).** The refusal of zipfile and its
  warning about two entries at one local header, outside and inside.
* **zip_traversal on the raw name (answer O5).** Read from the name as the
  directory carries it, before any member is read, never from the cleaned
  display name.
* **What stays a limit (answer O7).** Depth, entry count and the declared sum
  keep archive_limit.
* **package_entries and xml_entities (finding F3).** An ODF or OOXML package
  with more than ARCHIVE_MAX_ENTRIES directory records is counted before
  ZipFile is ever called; an XML part whose parser reports entity
  amplification is the entity bomb. Both rules are proven with real
  exceptions of lxml and expat, so a changed library message turns a test red
  instead of falling back to xml_invalid unnoticed.
* **image_pixels and mail_nesting (plan 36-09, answer O6).** A picture Pillow
  refuses (over 100 MP) that declares more than 1000 pixels per byte of its
  file, read from the header alone; a mail whose message/rfc822 parts nest
  deeper than 32 levels, measured on the parsed mail or, when the parser runs
  into RecursionError, counted over the raw bytes.
* **cfb_cycle (plan 36-09).** A cycle in the FAT chain of a stream or of the
  mini stream, in the MiniFAT, in the DIFAT chain, in the chain of directory
  sectors and in the directory tree, for .doc and .xls; a structural error
  without a cycle stays corrupt, and the sniff keeps its answers.
"""

from __future__ import annotations

import inspect
import io
import logging
import warnings
import xml.etree.ElementTree as ElementTree
import zipfile
from collections.abc import Callable, Iterator
from email.message import EmailMessage
from pathlib import Path
from typing import Any

import openpyxl
import pptx
import pytest
from lxml import etree  # pyright: ignore[reportAttributeAccessIssue]
from PIL import Image

from findling import config
from findling.config import (
    ARCHIVE_MALICIOUS_RATIO,
    ARCHIVE_MAX_ENTRIES,
    ARCHIVE_MAX_RATIO,
    ARCHIVE_RATIO_MIN_BYTES,
    settings,
)
from findling.extract import archive, cfb, image, legacy_doc, legacy_xls, mail, odf, office
from findling.extract.archive import extract_zip
from findling.extract.errors import ExtractionOutcome, Reason, State, is_entity_bomb
from hostile import (
    ATTACK_RATIO,
    DOCX_DOCUMENT,
    ODT_CONTENT,
    PPTX_SLIDE,
    XLSX_SHEET,
    Layout,
    Member,
    _difat_cycle,
    _directory_cycle,
    _fat_cycle,
    _fat_declared_far_past_the_file,
    _minifat_cycle,
    _png_chunk,
    _same_offset_twice,
    billion_laughs_xml,
    build_doc,
    build_workbook,
    build_zip,
    docx_with,
    honest_large_jpeg,
    nested_rfc822,
    odt_with,
    package_with_entries,
    png_bomb,
    pptx_with,
    quadratic_blowup_xml,
    text_member,
    xlsx_with,
    xml_text,
    zip_depth_three,
    zip_entries,
    zip_moderate_ratio_member,
    zip_overlap,
    zip_ratio_member,
    zip_traversal,
)


@pytest.fixture(autouse=True)
def _fresh_settings() -> Iterator[None]:
    settings.cache_clear()
    yield
    settings.cache_clear()


def _run(tmp_path: Path, raw: bytes) -> tuple[ExtractionOutcome, list[str]]:
    """extract_zip in this process, and what stands next to the job file afterwards."""
    target = tmp_path / "job-8.part"
    target.write_bytes(raw)
    outcome = extract_zip(str(target))
    return outcome, sorted(path.name for path in tmp_path.iterdir())


def _blocked(outcome: ExtractionOutcome, detail: str) -> None:
    assert (outcome.state, outcome.reason, outcome.detail) == (State.SKIPPED, Reason.BLOCKED_MALICIOUS, detail)
    assert outcome.text == ""


def _limit(outcome: ExtractionOutcome) -> None:
    assert (outcome.state, outcome.reason, outcome.detail) == (State.SKIPPED, Reason.ARCHIVE_LIMIT, None)


def _ratio_of(raw: bytes, name: str) -> float:
    """The ratio archive.py judges, read back from the directory of the bytes."""
    with zipfile.ZipFile(io.BytesIO(raw)) as check:
        info = check.getinfo(name)
    return info.file_size / max(info.compress_size, 1)


def _declared(packed: int, declared: int) -> bytes:
    """A stored member of ``packed`` bytes that declares ``declared`` bytes, next to a plain member."""
    return build_zip([text_member("vornwort"), Member(b"dicht.bin", b"\x01" * packed, declared=declared)])


# -- the owner's number ------------------------------------------------------------


def test_the_corpus_and_the_code_carry_the_same_owner_number() -> None:
    assert ATTACK_RATIO == config.ARCHIVE_MALICIOUS_RATIO == 1000
    assert ARCHIVE_MAX_RATIO < ARCHIVE_MALICIOUS_RATIO


# -- zip_ratio, two steps (O12) ------------------------------------------------------


def test_a_member_just_over_the_first_ratio_is_an_archive_limit(tmp_path: Path) -> None:
    raw = _declared(10_383, 101 * 10_383)
    assert _ratio_of(raw, "dicht.bin") == 101
    with zipfile.ZipFile(io.BytesIO(raw)) as check:
        assert check.getinfo("dicht.bin").file_size >= ARCHIVE_RATIO_MIN_BYTES + 1

    _limit(_run(tmp_path, raw)[0])


def test_a_member_of_about_150_to_1_is_an_archive_limit(tmp_path: Path) -> None:
    raw = zip_moderate_ratio_member()
    assert ARCHIVE_MAX_RATIO < _ratio_of(raw, "dicht.bin") < 200

    _limit(_run(tmp_path, raw)[0])


def test_a_member_just_under_the_attack_ratio_of_deflate_is_an_archive_limit(tmp_path: Path) -> None:
    raw = zip_moderate_ratio_member(950)
    assert 900 < _ratio_of(raw, "dicht.bin") <= ARCHIVE_MALICIOUS_RATIO

    _limit(_run(tmp_path, raw)[0])


def test_a_member_exactly_at_the_attack_ratio_is_an_archive_limit(tmp_path: Path) -> None:
    raw = _declared(1_049, ARCHIVE_MALICIOUS_RATIO * 1_049)
    assert _ratio_of(raw, "dicht.bin") == ARCHIVE_MALICIOUS_RATIO

    _limit(_run(tmp_path, raw)[0])


def test_a_member_one_byte_over_the_attack_ratio_is_blocked(tmp_path: Path) -> None:
    raw = _declared(1_049, ARCHIVE_MALICIOUS_RATIO * 1_049 + 1)
    assert _ratio_of(raw, "dicht.bin") > ARCHIVE_MALICIOUS_RATIO

    outcome, left = _run(tmp_path, raw)

    _blocked(outcome, "zip_ratio")
    assert left == ["job-8.part"]


def test_two_mib_of_zeros_under_deflate_are_blocked(tmp_path: Path) -> None:
    raw = zip_ratio_member()
    assert _ratio_of(raw, "nullen.bin") > ARCHIVE_MALICIOUS_RATIO

    _blocked(_run(tmp_path, raw)[0], "zip_ratio")


def test_the_attack_ratio_wins_over_a_member_between_the_ratios(tmp_path: Path) -> None:
    raw = build_zip(
        [
            Member(b"dicht.bin", b"\x01" * 10_383, declared=101 * 10_383),
            Member(b"bombe.bin", b"\x01" * 1_049, declared=ARCHIVE_MALICIOUS_RATIO * 1_049 + 1),
        ]
    )

    _blocked(_run(tmp_path, raw)[0], "zip_ratio")


def test_a_small_member_is_not_judged_by_its_ratio(tmp_path: Path) -> None:
    raw = _declared(1_048, ARCHIVE_RATIO_MIN_BYTES - 1)
    assert _ratio_of(raw, "dicht.bin") > ARCHIVE_MALICIOUS_RATIO

    outcome, _ = _run(tmp_path, raw)

    assert outcome.state is State.INDEXED
    assert "vornwort" in outcome.text


def test_a_member_exactly_at_the_first_ratio_is_not_judged(tmp_path: Path) -> None:
    raw = _declared(10_486, ARCHIVE_MAX_RATIO * 10_486)
    assert _ratio_of(raw, "dicht.bin") == ARCHIVE_MAX_RATIO
    with zipfile.ZipFile(io.BytesIO(raw)) as check:
        assert check.getinfo("dicht.bin").file_size >= ARCHIVE_RATIO_MIN_BYTES

    outcome, _ = _run(tmp_path, raw)

    assert outcome.state is State.INDEXED
    assert "vornwort" in outcome.text


def test_the_attack_ratio_is_read_through_the_module(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    raw = zip_ratio_member()
    monkeypatch.setattr(archive, "ARCHIVE_MALICIOUS_RATIO", 10**9)

    _limit(_run(tmp_path, raw)[0])


def test_the_attack_ratio_counts_inside_an_inner_zip_file(tmp_path: Path) -> None:
    inner = _declared(1_049, ARCHIVE_MALICIOUS_RATIO * 1_049 + 1)
    raw = build_zip([text_member("aussenwort"), Member(b"innen.zip", inner)])

    _blocked(_run(tmp_path, raw)[0], "zip_ratio")


# -- zip_overlap on every level (O2) --------------------------------------------------


def test_overlapping_entries_outside_are_blocked(tmp_path: Path) -> None:
    raw = zip_overlap()
    with zipfile.ZipFile(io.BytesIO(raw)) as check, pytest.raises(zipfile.BadZipFile, match="Overlapped"):
        check.read("a.txt")

    outcome, left = _run(tmp_path, raw)

    _blocked(outcome, "zip_overlap")
    assert left == ["job-8.part"]


def test_overlapping_entries_inside_are_blocked(tmp_path: Path) -> None:
    raw = build_zip([Member(b"innen.zip", zip_overlap()), text_member("aussenwort")])

    _blocked(_run(tmp_path, raw)[0], "zip_overlap")


def test_two_entries_at_one_local_header_are_blocked_outside(tmp_path: Path) -> None:
    raw = _same_offset_twice(b"o.txt", b"overlapwort")
    # zipfile warns instead of refusing, for one of the two entries.
    with zipfile.ZipFile(io.BytesIO(raw)) as check, warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        for info in check.infolist():
            check.read(info)
    assert [str(warning.message).startswith("Overlapped entries") for warning in caught] == [True]

    _blocked(_run(tmp_path, raw)[0], "zip_overlap")


def test_two_entries_at_one_local_header_are_blocked_inside(tmp_path: Path) -> None:
    raw = build_zip([Member(b"innen.zip", _same_offset_twice(b"o.txt", b"overlapwort")), text_member("aussenwort")])

    _blocked(_run(tmp_path, raw)[0], "zip_overlap")


def test_duplicate_names_at_their_own_headers_are_no_overlap(tmp_path: Path) -> None:
    raw = build_zip([Member(b"gleich.txt", b"erstwort"), Member(b"gleich.txt", b"zweitwort")])

    outcome, _ = _run(tmp_path, raw)

    assert outcome.state is State.INDEXED
    assert "erstwort" in outcome.text
    assert "zweitwort" in outcome.text


# -- zip_traversal on the raw name (O5) -----------------------------------------------

_CLIMBING = (
    b"../x.txt",
    b"a/../../b",
    b"..\\x",
    b"/etc/x",
    b"\\x",
    b"C:\\x",
    b"c:x",
    b"a\x00b",
    b"a/..",
    b"..",
)
_NOT_CLIMBING = (b"..x.txt", b"a..b/c.txt", b"x/..y", b"x../y.txt", b"...txt", b"a/b:c.txt")


@pytest.mark.parametrize("name", _CLIMBING, ids=repr)
def test_a_name_that_climbs_out_is_blocked(name: bytes, tmp_path: Path) -> None:
    outcome, left = _run(tmp_path, zip_traversal(name))

    _blocked(outcome, "zip_traversal")
    assert left == ["job-8.part"]


@pytest.mark.parametrize("name", _CLIMBING, ids=repr)
def test_the_rule_flags_every_climbing_name(name: bytes) -> None:
    assert archive._is_traversal(name.decode("ascii"))


@pytest.mark.parametrize("name", _NOT_CLIMBING, ids=repr)
def test_a_name_with_dots_that_do_not_climb_stays_text(name: bytes, tmp_path: Path) -> None:
    assert not archive._is_traversal(name.decode("ascii"))

    outcome, _ = _run(tmp_path, zip_traversal(name))

    assert outcome.state is State.INDEXED
    assert "aussenwort" in outcome.text


def test_a_climbing_name_inside_an_inner_zip_file_is_blocked(tmp_path: Path) -> None:
    raw = build_zip([text_member("aussenwort"), Member(b"innen.zip", zip_traversal(b"../x.txt"))])

    _blocked(_run(tmp_path, raw)[0], "zip_traversal")


def test_a_climbing_name_is_judged_before_any_member_is_read(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    reads: list[str] = []
    real = archive._read_member

    def counting(*args: object, **kwargs: object) -> None:
        reads.append("member")
        real(*args, **kwargs)  # pyright: ignore[reportArgumentType]

    monkeypatch.setattr(archive, "_read_member", counting)
    # The climbing name is the last member: the plain ones in front of it are not read.
    raw = build_zip([text_member("vornwort"), text_member("mittewort"), Member(b"../x.txt", b"x")])

    _blocked(_run(tmp_path, raw)[0], "zip_traversal")
    assert reads == []


def test_the_rule_reads_the_raw_name_and_never_the_display_name() -> None:
    rule = inspect.getsource(archive._is_traversal)
    caller = inspect.getsource(archive._read_zip)

    assert "clean_display_name" not in rule
    assert "_is_traversal(info.orig_filename)" in caller
    assert "clean_display_name" not in caller
    # zipfile cuts a name at NUL in ZipInfo.filename; the raw name keeps it.
    info = zipfile.ZipInfo("a\x00b")
    assert info.filename == "a"
    assert info.orig_filename == "a\x00b"


# -- what stays a limit (O7) ---------------------------------------------------------


def test_depth_three_without_zeros_is_an_archive_limit(tmp_path: Path) -> None:
    _limit(_run(tmp_path, zip_depth_three())[0])


def test_10001_entries_are_an_archive_limit(tmp_path: Path) -> None:
    _limit(_run(tmp_path, zip_entries())[0])


def test_a_lying_huge_declaration_is_an_archive_limit_before_the_ratio(tmp_path: Path) -> None:
    raw = build_zip([Member(b"a.txt", b"hallo", declared=0xFFFFFFF0)])
    assert _ratio_of(raw, "a.txt") > ARCHIVE_MALICIOUS_RATIO

    _limit(_run(tmp_path, raw)[0])


def test_an_over_limit_without_a_signature_carries_no_code() -> None:
    assert archive.OverLimit().signature is None
    assert archive.OverLimit("zip_ratio").signature == "zip_ratio"


# -- package_entries (F3) ---------------------------------------------------------------

_READERS: dict[str, Callable[[str], ExtractionOutcome]] = {
    "docx": office.extract_docx,
    "pptx": office.extract_pptx,
    "xlsx": office.extract_xlsx,
    "odt": odf.extract_odf,
}


def _judged(read: Callable[[str], ExtractionOutcome], path: Path) -> ExtractionOutcome:
    """A reader as the child calls it: what escapes is translated at the boundary."""
    try:
        return read(str(path))
    except Exception as error:  # the boundary of the sandbox, in this process
        return ExtractionOutcome.from_exception(error)


def _package(tmp_path: Path, kind: str, count: int) -> Path:
    # package_with_entries knows the first part of docx, xlsx and odt; a pptx
    # names the same [Content_Types].xml as the other two OOXML formats.
    path = tmp_path / f"paket-{kind}.part"
    path.write_bytes(package_with_entries("docx" if kind == "pptx" else kind, count))
    return path


@pytest.fixture
def no_zipfile(monkeypatch: pytest.MonkeyPatch) -> list[str]:
    calls: list[str] = []

    def refuse(*args: Any, **kwargs: Any) -> zipfile.ZipFile:
        calls.append("ZipFile")
        raise AssertionError("ZipFile must not be reached")

    for module in (odf, office, zipfile):
        monkeypatch.setattr(module, "ZipFile", refuse)
    return calls


@pytest.mark.parametrize("kind", list(_READERS))
def test_a_package_of_10001_records_is_blocked_before_zipfile(kind: str, tmp_path: Path, no_zipfile: list[str]) -> None:
    path = _package(tmp_path, kind, ARCHIVE_MAX_ENTRIES + 1)

    outcome = _READERS[kind](str(path))

    _blocked(outcome, "package_entries")
    assert no_zipfile == []


@pytest.mark.parametrize("kind", list(_READERS))
def test_a_package_of_a_million_records_is_blocked_before_zipfile(
    kind: str, tmp_path: Path, no_zipfile: list[str]
) -> None:
    # The directory of a million records is longer than
    # ARCHIVE_MAX_CENTRAL_DIRECTORY_BYTES: the count ends at its length.
    path = _package(tmp_path, kind, 1_000_000)

    _blocked(_READERS[kind](str(path)), "package_entries")
    assert no_zipfile == []


@pytest.mark.parametrize("kind", list(_READERS))
def test_a_package_of_10000_records_is_left_to_the_loader(kind: str, tmp_path: Path) -> None:
    outcome = _judged(_READERS[kind], _package(tmp_path, kind, ARCHIVE_MAX_ENTRIES))

    assert outcome.reason is not Reason.BLOCKED_MALICIOUS
    assert outcome.state is State.FAILED


@pytest.mark.parametrize("kind", list(_READERS))
def test_a_directory_the_count_calls_corrupt_is_left_to_the_loader(kind: str, tmp_path: Path) -> None:
    path = tmp_path / "kaputt.part"
    path.write_bytes(b"PK\x03\x04 kein Verzeichnis")
    with path.open("rb") as handle, pytest.raises(archive.CorruptDirectory):
        archive.count_central_directory(handle, path.stat().st_size, limit=ARCHIVE_MAX_ENTRIES)

    outcome = _judged(_READERS[kind], path)

    assert (outcome.state, outcome.reason) == (State.FAILED, Reason.CORRUPT)


def test_every_reader_counts_before_it_opens_the_package() -> None:
    for module in (odf, office):
        count = inspect.getsource(module._too_many_entries)
        assert "count_central_directory" in count
        assert "ZipFile(" not in count
    for source in (inspect.getsource(odf.extract_odf), inspect.getsource(office._refused_before_the_loader)):
        assert source.index("_too_many_entries(path)") < source.index("ZipFile(")
    for reader in (office.extract_docx, office.extract_pptx, office.extract_xlsx):
        source = inspect.getsource(reader)
        assert "_refused_before_the_loader(path" in source
        assert "ZipFile(" not in source


# -- xml_entities (F3) -------------------------------------------------------------------

_WRAPS: dict[str, tuple[Callable[[bytes | None], bytes], bytes]] = {
    "odt": (odt_with, ODT_CONTENT),
    "docx": (docx_with, DOCX_DOCUMENT),
    "pptx": (pptx_with, PPTX_SLIDE),
    "xlsx": (xlsx_with, XLSX_SHEET),
}
_BOMBS = {"billion-laughs": billion_laughs_xml, "quadratic-blowup": quadratic_blowup_xml}


@pytest.mark.parametrize("bomb", list(_BOMBS))
@pytest.mark.parametrize("kind", list(_WRAPS))
def test_an_entity_bomb_in_a_package_is_blocked(
    kind: str, bomb: str, tmp_path: Path, caplog: pytest.LogCaptureFixture
) -> None:
    caplog.set_level(logging.DEBUG)
    wrap, template = _WRAPS[kind]
    path = tmp_path / f"bombe-{kind}.part"
    path.write_bytes(wrap(_BOMBS[bomb](template=template)))

    outcome = _READERS[kind](str(path))

    _blocked(outcome, "xml_entities")
    assert "amplification" not in repr(outcome)
    assert "amplification" not in caplog.text


@pytest.mark.parametrize("kind", list(_WRAPS))
def test_an_ordinary_syntax_error_keeps_its_verdict(kind: str, tmp_path: Path) -> None:
    wrap, template = _WRAPS[kind]
    path = tmp_path / f"kaputt-{kind}.part"
    path.write_bytes(wrap(xml_text(b"<offen", template)))

    outcome = _judged(_READERS[kind], path)

    # lxml says xml_invalid, expat under openpyxl corrupt: the verdicts the
    # corpus of 36-07 measured before this plan.
    expected = Reason.CORRUPT if kind == "xlsx" else Reason.XML_INVALID
    assert (outcome.state, outcome.reason) == (State.FAILED, expected)


def test_lxml_reports_amplification_and_the_rule_sees_it() -> None:
    parser = etree.XMLParser(resolve_entities=False, no_network=True, load_dtd=False)
    with pytest.raises(etree.XMLSyntaxError) as raised:
        etree.fromstring(billion_laughs_xml(), parser=parser)

    assert is_entity_bomb(raised.value)


def test_expat_reports_amplification_and_the_rule_sees_it() -> None:
    with pytest.raises(ElementTree.ParseError) as raised:
        ElementTree.fromstring(billion_laughs_xml())  # noqa: S314 - the real expat exception is the point

    assert is_entity_bomb(raised.value)


def test_an_ordinary_syntax_error_is_no_entity_bomb() -> None:
    with pytest.raises(etree.XMLSyntaxError) as of_lxml:
        etree.fromstring(b"<r><offen</r>")
    with pytest.raises(ElementTree.ParseError) as of_expat:
        ElementTree.fromstring(b"<r><offen</r>")  # noqa: S314 - the real expat exception is the point

    assert not is_entity_bomb(of_lxml.value)
    assert not is_entity_bomb(of_expat.value)
    assert not is_entity_bomb(ValueError("nothing of the kind"))


# -- the counter-check: real packages written by their own libraries -------------------


def test_a_large_real_xlsx_stays_indexed(tmp_path: Path) -> None:
    book = openpyxl.Workbook()
    sheet = book.active
    assert sheet is not None
    for row in range(5_000):
        sheet.append([f"zeile{row}", row, row * 0.5, "tabellenwort"])
    for number in range(20):
        book.create_sheet(f"blatt{number}").append(["blattwort", number])
    path = tmp_path / "gross.part"
    book.save(path)

    outcome = office.extract_xlsx(str(path))

    assert outcome.state is State.INDEXED
    assert "tabellenwort" in outcome.text


def test_a_large_real_pptx_stays_indexed(tmp_path: Path) -> None:
    deck = pptx.Presentation()
    for number in range(200):
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        title = slide.shapes.title
        assert title is not None
        title.text = f"folie{number} foliewort"
    path = tmp_path / "gross.part"
    deck.save(str(path))

    outcome = office.extract_pptx(str(path))

    assert outcome.state is State.INDEXED
    assert "folie199" in outcome.text


# -- image_pixels (plan 36-09, answer O6) ----------------------------------------------


def _picture(tmp_path: Path, raw: bytes, name: str = "bild.part") -> str:
    path = tmp_path / name
    path.write_bytes(raw)
    return str(path)


def _png_of_size(width: int, height: int, size: int) -> bytes:
    """A PNG bomb padded with an unknown ancillary chunk to exactly ``size`` bytes."""
    bare = png_bomb(width, height)
    room = size - len(bare) - 12  # length, type and CRC of the padding chunk
    assert room >= 0
    # Behind the signature (8 bytes) and IHDR (25 bytes); Pillow skips a chunk it does not know.
    raw = bare[:33] + _png_chunk(b"prVt", bytes(room)) + bare[33:]
    assert len(raw) == size
    return raw


def test_the_bomb_number_of_the_code_is_the_owner_number() -> None:
    assert image.IMAGE_BOMB_PIXELS_PER_BYTE == 1000
    assert mail.MAIL_MAX_NESTING == 32


def test_a_png_bomb_is_blocked_without_a_row_decoded(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    loads: list[object] = []
    monkeypatch.setattr(Image.Image, "load", lambda self: loads.append(self))
    before = Image.MAX_IMAGE_PIXELS

    outcome = image.extract_image(_picture(tmp_path, png_bomb()))

    _blocked(outcome, "image_pixels")
    assert loads == []
    assert before == Image.MAX_IMAGE_PIXELS


def test_an_honest_large_jpeg_stays_too_large(tmp_path: Path) -> None:
    raw = honest_large_jpeg()
    assert image.IMAGE_BOMB_PIXELS_PER_BYTE * len(raw) > 12_000 * 8_600

    outcome = image.extract_image(_picture(tmp_path, raw))

    assert (outcome.state, outcome.reason, outcome.detail) == (State.SKIPPED, Reason.TOO_LARGE, None)


def test_exactly_1000_pixels_per_byte_stay_too_large_and_one_byte_less_is_blocked(tmp_path: Path) -> None:
    # 20 000 x 20 000 = 400 MP over Pillow's doubled cap; 400 000 bytes are exactly 1000 pixels a byte.
    at = image.extract_image(_picture(tmp_path, _png_of_size(20_000, 20_000, 400_000), "an.part"))
    over = image.extract_image(_picture(tmp_path, _png_of_size(20_000, 20_000, 399_999), "drueber.part"))

    assert (at.state, at.reason) == (State.SKIPPED, Reason.TOO_LARGE)
    _blocked(over, "image_pixels")


def test_a_picture_pillow_only_warns_about_stays_too_large(tmp_path: Path) -> None:
    # 60 MP in a few hundred bytes: over the cap of image.py, under twice it, so
    # Pillow does not refuse it and the bomb rule is never asked.
    with pytest.warns(Image.DecompressionBombWarning):
        outcome = image.extract_image(_picture(tmp_path, png_bomb(8000, 7500)))

    assert (outcome.state, outcome.reason, outcome.detail) == (State.SKIPPED, Reason.TOO_LARGE, None)


def test_the_pixel_cap_is_put_back_when_the_second_open_fails(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    real_open = Image.open
    seen: list[object] = []

    def second_fails(*args: Any, **kwargs: Any) -> Any:
        seen.append(Image.MAX_IMAGE_PIXELS)
        if len(seen) == 2:
            raise OSError("second look")
        return real_open(*args, **kwargs)

    monkeypatch.setattr(image.Image, "open", second_fails)
    before = Image.MAX_IMAGE_PIXELS

    outcome = image.extract_image(_picture(tmp_path, png_bomb()))

    assert (outcome.state, outcome.reason, outcome.detail) == (State.SKIPPED, Reason.TOO_LARGE, None)
    assert seen == [before, None]
    assert before == Image.MAX_IMAGE_PIXELS


def test_the_pixel_cap_is_put_back_after_the_bomb(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    real_open = Image.open
    seen: list[object] = []

    def recording(*args: Any, **kwargs: Any) -> Any:
        seen.append(Image.MAX_IMAGE_PIXELS)
        return real_open(*args, **kwargs)

    monkeypatch.setattr(image.Image, "open", recording)
    before = Image.MAX_IMAGE_PIXELS

    _blocked(image.extract_image(_picture(tmp_path, png_bomb())), "image_pixels")
    assert seen == [before, None]
    assert before == Image.MAX_IMAGE_PIXELS


# -- mail_nesting (plan 36-09, answer O6) ------------------------------------------------


def _mail(tmp_path: Path, raw: bytes) -> ExtractionOutcome:
    path = tmp_path / "post.part"
    path.write_bytes(raw)
    return mail.extract_eml(str(path))


def test_33_forwarded_levels_are_blocked(tmp_path: Path) -> None:
    _blocked(_mail(tmp_path, nested_rfc822(33)), "mail_nesting")


def test_32_forwarded_levels_are_read(tmp_path: Path) -> None:
    # The outer subject is read; the forwarded mails are attachments, read by name only (D-31-04).
    outcome = _mail(tmp_path, nested_rfc822(32))

    assert (outcome.state, outcome.reason) == (State.INDEXED, None)
    assert outcome.text == "s"


def test_1000_forwarded_levels_are_blocked_and_not_corrupt(tmp_path: Path) -> None:
    _blocked(_mail(tmp_path, nested_rfc822(1000)), "mail_nesting")


def test_the_depth_is_measured_on_the_parsed_mail() -> None:
    from email import message_from_bytes

    for levels in (0, 1, 32, 33):
        assert mail.forwarding_depth(message_from_bytes(nested_rfc822(levels))) == levels


def test_40_forwarded_mails_side_by_side_are_read(tmp_path: Path) -> None:
    outer = EmailMessage()
    outer["Subject"] = "Sammelpost"
    outer.set_content("Anbei vierzig Mails.")
    for number in range(40):
        inner = EmailMessage()
        inner["Subject"] = f"Mail {number}"
        inner.set_content(f"geschwisterwort {number}")
        outer.add_attachment(inner)
    raw = outer.as_bytes()
    assert mail.count_forward_headers(raw) == 40

    outcome = _mail(tmp_path, raw)

    assert outcome.state is State.INDEXED
    assert "Sammelpost" in outcome.text


def test_the_raw_count_reads_every_spelling() -> None:
    raw = b"content-type: MESSAGE/RFC822\r\n\r\nContent-Type:message/rfc822\rCONTENT-TYPE: \tMessage/Rfc822\n"
    assert mail.count_forward_headers(raw) == 3
    assert mail.count_forward_headers(b"X-Content-Type: message/rfc822\r\n") == 0


def _raise_recursion(*args: object, **kwargs: object) -> None:
    raise RecursionError


def test_a_recursion_error_over_many_forward_headers_is_mail_nesting(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(mail.BytesParser, "parse", _raise_recursion)

    _blocked(_mail(tmp_path, nested_rfc822(33)), "mail_nesting")


def test_a_recursion_error_under_few_forward_headers_stays_corrupt(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(mail.BytesParser, "parse", _raise_recursion)

    outcome = _mail(tmp_path, nested_rfc822(32))

    assert outcome == ExtractionOutcome.failed(Reason.CORRUPT, detail="RecursionError")


# -- cfb_cycle (plan 36-09) -------------------------------------------------------------


def _container_cycle(layout: Layout) -> None:
    """The FAT chain of the mini stream container links its first sector to itself."""
    start = layout.entries[0][5]
    layout.fat[start] = start


def _directory_chain_cycle(layout: Layout) -> None:
    """The FAT chain of the directory sectors links its first sector to itself."""
    start = layout.header["dir_start"]
    layout.fat[start] = start


_DOC_TEXT = b"olewort im text\r"

_CYCLES: dict[str, Callable[[Layout], None]] = {
    "fat": _fat_cycle,
    "fat-of-the-mini-stream": _container_cycle,
    "minifat": _minifat_cycle,
    "difat": _difat_cycle,
    "directory-chain": _directory_chain_cycle,
    "directory-tree": _directory_cycle,
}

# The Workbook of build_workbook is small and lives in the mini stream, so the
# FAT case of a workbook is the chain of the mini stream container.
_XLS_CYCLES = {name: edit for name, edit in _CYCLES.items() if name != "fat"}


def _ole(tmp_path: Path, raw: bytes, name: str = "ole.part") -> str:
    path = tmp_path / name
    path.write_bytes(raw)
    return str(path)


@pytest.mark.parametrize("cycle", list(_CYCLES))
def test_every_cycle_of_a_word_file_is_blocked(cycle: str, tmp_path: Path) -> None:
    _blocked(legacy_doc.extract_doc(_ole(tmp_path, build_doc(_DOC_TEXT, edit=_CYCLES[cycle]))), "cfb_cycle")


@pytest.mark.parametrize("cycle", list(_XLS_CYCLES))
def test_every_cycle_of_an_excel_file_is_blocked(cycle: str, tmp_path: Path) -> None:
    raw = build_workbook([(0, 0, 0)], ["zellwort"], edit=_XLS_CYCLES[cycle])

    _blocked(legacy_xls.extract_xls(_ole(tmp_path, raw)), "cfb_cycle")


def test_the_untouched_files_are_read(tmp_path: Path) -> None:
    assert "olewort" in legacy_doc.extract_doc(_ole(tmp_path, build_doc(_DOC_TEXT), "a.part")).text
    assert (
        "zellwort" in legacy_xls.extract_xls(_ole(tmp_path, build_workbook([(0, 0, 0)], ["zellwort"]), "b.part")).text
    )


def _fat_past_the_file(layout: Layout) -> None:
    """A link of the WordDocument chain to a sector behind the end of the file: no cycle."""
    start = layout.entries[1][5]
    layout.fat[start + 2] = 100


@pytest.mark.parametrize("edit", [_fat_declared_far_past_the_file, _fat_past_the_file], ids=["fat-count", "link"])
def test_a_structural_error_without_a_cycle_stays_corrupt(edit: Callable[[Layout], None], tmp_path: Path) -> None:
    outcome = legacy_doc.extract_doc(_ole(tmp_path, build_doc(_DOC_TEXT, edit=edit)))

    assert (outcome.state, outcome.reason) == (State.FAILED, Reason.CORRUPT)


def test_a_cycle_is_a_corrupt_compound_file_for_every_caller_of_the_parent() -> None:
    assert issubclass(cfb.CompoundCycle, cfb.CorruptCompoundFile)


def test_the_cycle_check_leaves_every_other_defect_to_olefile(tmp_path: Path) -> None:
    # No cycle, so no raise, whatever else is wrong with the file.
    for name, raw in {
        "fat-count": build_doc(_DOC_TEXT, edit=_fat_declared_far_past_the_file),
        "link": build_doc(_DOC_TEXT, edit=_fat_past_the_file),
        "short": build_doc(_DOC_TEXT)[:700],
        "no-ole": b"PK\x03\x04" + bytes(600),
    }.items():
        cfb.check_cycles(_ole(tmp_path, raw, f"{name}.part"))


@pytest.mark.parametrize(
    ("edit", "kind"),
    [(_directory_cycle, "word"), (_directory_chain_cycle, None), (_fat_cycle, "word")],
    ids=["tree", "directory-chain", "fat"],
)
def test_the_sniff_keeps_its_answer_for_a_cycle(
    edit: Callable[[Layout], None], kind: str | None, tmp_path: Path
) -> None:
    # A cycle ends the walk of the sniff and not the sniff (cfb.py): the tree
    # walk skips what it has seen, the directory chain ends in None as before.
    assert cfb.ole_kind(_ole(tmp_path, build_doc(_DOC_TEXT, edit=edit))) == kind
